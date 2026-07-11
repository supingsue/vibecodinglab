# -*- coding: utf-8 -*-
"""pdf2pptx.py 的 pytest 測試。

測試用的 3 頁 PDF 以 PyMuPDF 動態建立（不依賴外部測試檔）：
第 1 頁含兩行文字、第 2 頁為純圖形（無文字）、第 3 頁含一行文字。
"""

import sys
from pathlib import Path

import fitz
import pytest
from pptx import Presentation

# 讓測試可以直接 import 專案根目錄的 pdf2pptx.py
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pdf2pptx  # noqa: E402


@pytest.fixture
def sample_pdf(tmp_path):
    """動態建立 3 頁測試 PDF，回傳檔案路徑。"""
    path = tmp_path / "sample.pdf"
    doc = fitz.open()
    # 第 1 頁：兩行文字
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), "Page one text", fontsize=14)
    page.insert_text((72, 130), "Second line", fontsize=14)
    # 第 2 頁：純圖形、無文字
    page = doc.new_page(width=595, height=842)
    page.draw_rect(fitz.Rect(100, 100, 400, 400), fill=(0, 1, 0))
    # 第 3 頁：一行文字
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), "Page three text", fontsize=14)
    doc.save(str(path))
    doc.close()
    return path


def notes_of(slide) -> str:
    """取出投影片的備忘稿文字；沒有備忘稿視為空字串。"""
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text


class TestPageCount:
    """頁數一致：轉出的投影片數應與 PDF 頁數相同。"""

    def test_all_pages_converted(self, sample_pdf, tmp_path):
        out = tmp_path / "out.pptx"
        assert pdf2pptx.main([str(sample_pdf), "-o", str(out)]) == 0
        prs = Presentation(str(out))
        assert len(prs.slides) == 3

    def test_default_output_path(self, sample_pdf):
        # 未指定 -o：輸出到同目錄、同檔名的 .pptx
        assert pdf2pptx.main([str(sample_pdf)]) == 0
        out = sample_pdf.with_suffix(".pptx")
        assert out.is_file()
        assert len(Presentation(str(out)).slides) == 3


class TestNotes:
    """備忘稿內容：文字頁含該頁文字（保留換行）、純圖片頁留空。"""

    def test_text_pages_have_notes(self, sample_pdf, tmp_path):
        out = tmp_path / "out.pptx"
        pdf2pptx.main([str(sample_pdf), "-o", str(out)])
        prs = Presentation(str(out))
        notes1 = notes_of(prs.slides[0])
        assert "Page one text" in notes1
        assert "Second line" in notes1
        assert "\n" in notes1  # 兩行文字之間的換行應保留
        assert "Page three text" in notes_of(prs.slides[2])

    def test_image_only_page_has_empty_notes(self, sample_pdf, tmp_path):
        out = tmp_path / "out.pptx"
        pdf2pptx.main([str(sample_pdf), "-o", str(out)])
        prs = Presentation(str(out))
        assert notes_of(prs.slides[1]).strip() == ""


class TestPagesOption:
    """--pages 頁數範圍：正常範圍、單頁、部分超界警告、格式錯誤。"""

    def test_range_produces_two_slides(self, sample_pdf, tmp_path):
        out = tmp_path / "out.pptx"
        assert pdf2pptx.main([str(sample_pdf), "-o", str(out),
                              "--pages", "2-3"]) == 0
        prs = Presentation(str(out))
        assert len(prs.slides) == 2
        # 第一張投影片對應原第 2 頁（純圖形頁），備忘稿應為空
        assert notes_of(prs.slides[0]).strip() == ""
        assert "Page three text" in notes_of(prs.slides[1])

    def test_single_page(self, sample_pdf, tmp_path):
        out = tmp_path / "out.pptx"
        assert pdf2pptx.main([str(sample_pdf), "-o", str(out),
                              "--pages", "3"]) == 0
        assert len(Presentation(str(out)).slides) == 1

    def test_partial_overrange_warns(self, sample_pdf, tmp_path, capsys):
        out = tmp_path / "out.pptx"
        assert pdf2pptx.main([str(sample_pdf), "-o", str(out),
                              "--pages", "2-20"]) == 0
        assert "文件僅有 3 頁，已轉換第 2-3 頁" in capsys.readouterr().out
        assert len(Presentation(str(out)).slides) == 2

    def test_full_overrange_exits_1(self, sample_pdf, tmp_path, capsys):
        out = tmp_path / "out.pptx"
        with pytest.raises(SystemExit) as excinfo:
            pdf2pptx.main([str(sample_pdf), "-o", str(out),
                           "--pages", "15-20"])
        assert excinfo.value.code == 1
        assert "超出範圍" in capsys.readouterr().err
        assert not out.exists()  # 不留下輸出檔

    @pytest.mark.parametrize("bad", ["5-2", "abc", "0-2", "1-2-3", "-3", ""])
    def test_format_error_exits_1(self, sample_pdf, tmp_path, capsys, bad):
        out = tmp_path / "out.pptx"
        assert pdf2pptx.main([str(sample_pdf), "-o", str(out),
                              "--pages", bad]) == 1
        assert "--pages 格式錯誤，範例：--pages 5-12" in capsys.readouterr().err
        assert not out.exists()  # 格式錯誤時不執行轉換


class TestFileNotFound:
    """檔案不存在：中文錯誤訊息、結束碼 1。"""

    def test_missing_input_exits_1(self, tmp_path, capsys):
        missing = tmp_path / "no_such.pdf"
        with pytest.raises(SystemExit) as excinfo:
            pdf2pptx.main([str(missing)])
        assert excinfo.value.code == 1
        assert f"找不到檔案：{missing}" in capsys.readouterr().err
